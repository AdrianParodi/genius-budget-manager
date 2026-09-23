"""
Budget Manager — Módulo de reportería
Extrae datos de la API REST y genera un archivo Excel con métricas de campañas.

Uso:
    python extract.py

Requisitos:
    pip install -r requirements.txt

El servidor de Budget Manager debe estar corriendo en http://localhost:8080
"""

import json
import urllib.request
import urllib.error
from datetime import datetime

BUDGET_MANAGER_URL = 'http://localhost:8080'
LANDING_CRM_URL    = 'http://localhost:3000'
OUTPUT_FILE        = 'report.xlsx'


def api_get(url: str) -> list | dict:
    with urllib.request.urlopen(url, timeout=5) as response:
        return json.loads(response.read())


def get_campaigns() -> list:
    """Obtiene todas las campañas del Budget Manager."""
    return api_get(f'{BUDGET_MANAGER_URL}/api/campaigns')


def get_budget_summary() -> dict:
    """Obtiene el resumen global de presupuesto (solo campañas activas)."""
    return api_get(f'{BUDGET_MANAGER_URL}/api/campaigns/summary')


def get_leads_summary() -> list:
    """Obtiene el resumen de leads por landing desde el Landing CRM."""
    return api_get(f'{LANDING_CRM_URL}/api/landings/summary')


def export_to_excel(campaigns: list, summary: dict) -> None:
    """Genera el archivo Excel con métricas de campañas."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    header_fill = PatternFill('solid', fgColor='17365D')
    header_font = Font(name='Aptos Display', size=11, bold=True, color='FFFFFF')
    body_font = Font(name='Aptos', size=10, color='1F1F1F')
    border = Border(
        left=Side(style='thin', color='D9E2F3'),
        right=Side(style='thin', color='D9E2F3'),
        top=Side(style='thin', color='D9E2F3'),
        bottom=Side(style='thin', color='D9E2F3'),
    )
    number_format = '#,##0.00'

    def style_sheet(ws, widths: dict, numeric_columns: list[int]) -> None:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.border = border
                cell.alignment = Alignment(vertical='center')
            for column_index in numeric_columns:
                row[column_index - 1].number_format = number_format

        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 24

    wb = openpyxl.Workbook()

    # Hoja de campañas
    ws = wb.active
    ws.title = 'Campañas'
    ws.append(['ID', 'Nombre', 'Cliente', 'Estado', 'Presupuesto', 'Gastado', 'Disponible'])
    for c in campaigns:
        ws.append([
            c.get('id'),
            c.get('name'),
            c.get('client'),
            c.get('status'),
            c.get('budget', 0),
            c.get('spent', 0),
            c.get('budget', 0) - c.get('spent', 0),
        ])

    # Hoja de resumen
    ws2 = wb.create_sheet('Resumen')
    ws2.append(['Métrica', 'Valor'])
    ws2.append(['Campañas activas',    summary.get('activeCampaigns', 0)])
    ws2.append(['Presupuesto total',   summary.get('totalBudget', 0)])
    ws2.append(['Total gastado',       summary.get('totalSpent', 0)])
    ws2.append(['Total disponible',    summary.get('totalAvailable', 0)])
    ws2.append(['% de consumo',        summary.get('consumptionPercentage', 0)])

    style_sheet(
        ws,
        {'A': 10, 'B': 28, 'C': 22, 'D': 14, 'E': 16, 'F': 16, 'G': 16},
        [5, 6, 7],
    )
    style_sheet(ws2, {'A': 24, 'B': 18}, [2])
    ws2['B6'].number_format = '0.00"%"'

    wb.save(OUTPUT_FILE)
    print(f'Reporte guardado en {OUTPUT_FILE}')


if __name__ == '__main__':
    print(f'Extrayendo datos — {datetime.now().strftime("%Y-%m-%d %H:%M")}')
    campaigns = get_campaigns()
    summary   = get_budget_summary()
    print(f'Campañas encontradas: {len(campaigns)}')
    export_to_excel(campaigns, summary)
